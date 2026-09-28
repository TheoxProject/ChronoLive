import io

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)


# ============================================================
# OUTILS
# ============================================================

def formater_temps(centisecondes):

    if centisecondes is None:
        return "-"

    heures = centisecondes // 360000

    minutes = (
        centisecondes // 6000
    ) % 60

    secondes = (
        centisecondes // 100
    ) % 60

    centiemes = (
        centisecondes % 100
    )

    return (
        f"{heures:02d}:"
        f"{minutes:02d}:"
        f"{secondes:02d}."
        f"{centiemes:02d}"
    )


def formater_ecart(centisecondes):

    if (
        centisecondes is None
        or centisecondes == 0
    ):
        return "-"

    return (
        "+"
        + formater_temps(
            centisecondes
        )
    )


# ============================================================
# EXPORT EXCEL
# ============================================================

def exporter_excel(
    resultat,
    titre,
    type_classement
):

    workbook = Workbook()

    feuille = workbook.active

    feuille.title = "Résultats"


    # ========================================================
    # COULEURS
    # ========================================================

    VERT = "5B8E73"
    VERT_PALE = "EDF4EF"
    VERT_TRES_PALE = "F7FAF8"
    GRIS = "E4EAE6"
    BLANC = "FFFFFF"
    TEXTE = "334155"


    # ========================================================
    # BORDURE
    # ========================================================

    bordure = Border(
        bottom=Side(
            style="thin",
            color=GRIS
        )
    )


    # ========================================================
    # TITRE
    # ========================================================

    feuille.merge_cells(
        "A1:G1"
    )


    cellule_titre = feuille["A1"]

    cellule_titre.value = titre

    cellule_titre.font = Font(
        name="Arial",
        size=18,
        bold=True,
        color=VERT
    )

    cellule_titre.alignment = Alignment(
        horizontal="left"
    )


    # ========================================================
    # EN-TETES
    # ========================================================

    entetes = [
        "Pos.",
        "Dossard",
        "Coureur",
        "Sexe",
        "Catégorie",
        "Temps",
        "Écart"
    ]


    for colonne, valeur in enumerate(
        entetes,
        start=1
    ):

        cellule = feuille.cell(
            row=3,
            column=colonne
        )

        cellule.value = valeur

        cellule.font = Font(
            name="Arial",
            size=10,
            bold=True,
            color=BLANC
        )

        cellule.fill = PatternFill(
            fill_type="solid",
            fgColor=VERT
        )

        cellule.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # ========================================================
    # DONNEES
    # ========================================================

    ligne = 4


    for index, coureur in enumerate(
        resultat
    ):

        if type_classement == "categorie_ffc":

            categorie = (
                coureur.get(
                    "categorie_ffc"
                )
                or "-"
            )

        elif type_classement == "scratch":

            categorie = (
                coureur.get(
                    "categorie"
                )
                or "-"
            )

            categorie_ffc = (
                coureur.get(
                    "categorie_ffc"
                )
                or "-"
            )

            categorie = (
                categorie
                + " / "
                + categorie_ffc
            )

        else:

            categorie = (
                coureur.get(
                    "categorie"
                )
                or "-"
            )


        dossard = (

            f"#{coureur['dossard']}"

            if coureur.get(
                "dossard"
            ) is not None

            else "-"

        )


        nom = (
            coureur.get(
                "nom",
                ""
            )
            + " "
            + coureur.get(
                "prenom",
                ""
            )
        ).strip()


        donnees = [

            coureur["position"],

            dossard,

            nom,

            coureur.get(
                "sexe",
                "-"
            ),

            categorie,

            formater_temps(
                coureur[
                    "temps_centisecondes"
                ]
            ),

            formater_ecart(
                coureur[
                    "ecart_centisecondes"
                ]
            )

        ]


        for colonne, valeur in enumerate(
            donnees,
            start=1
        ):

            cellule = feuille.cell(
                row=ligne,
                column=colonne
            )

            cellule.value = valeur

            cellule.font = Font(
                name="Arial",
                size=10,
                color=TEXTE
            )

            cellule.border = bordure

            cellule.alignment = Alignment(
                horizontal=(
                    "left"
                    if colonne == 3
                    else "center"
                ),
                vertical="center"
            )


            if index == 0:

                cellule.fill = PatternFill(
                    fill_type="solid",
                    fgColor=VERT_PALE
                )

            elif index % 2 == 1:

                cellule.fill = PatternFill(
                    fill_type="solid",
                    fgColor=VERT_TRES_PALE
                )

            else:

                cellule.fill = PatternFill(
                    fill_type=None
                )


        ligne += 1


    # ========================================================
    # LARGEUR DES COLONNES
    # ========================================================

    largeurs = {

        "A": 8,

        "B": 12,

        "C": 32,

        "D": 10,

        "E": 24,

        "F": 16,

        "G": 16

    }


    for colonne, largeur in (
        largeurs.items()
    ):

        feuille.column_dimensions[
            colonne
        ].width = largeur


    # ========================================================
    # HAUTEUR
    # ========================================================

    feuille.row_dimensions[1].height = 28

    feuille.row_dimensions[3].height = 24


    # ========================================================
    # GEL DES EN-TETES
    # ========================================================

    feuille.freeze_panes = "A4"


    # ========================================================
    # FILTRE
    # ========================================================

    if resultat:

        feuille.auto_filter.ref = (

            f"A3:G"
            f"{3 + len(resultat)}"

        )


    # ========================================================
    # SORTIE
    # ========================================================

    sortie = io.BytesIO()

    workbook.save(
        sortie
    )

    sortie.seek(0)

    return sortie.getvalue()

